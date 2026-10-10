from pathlib import Path
import docx

def extract_text_from_docx(file_path: Path | str) -> str:
    """
    Extracts text from a DOCX file, including both paragraphs and embedded tables.
    Handles corrupt, empty, and non-valid DOCX archives.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        raise ValueError(f"Corrupted or invalid DOCX document: {str(e)}")

    extracted_lines = []

    # 1. Paragraphs
    for para in doc.paragraphs:
        line = para.text.strip()
        if line:
            extracted_lines.append(line)

    # 2. Tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate repeated cell contents from merged cells
                unique_cells = []
                for cell in row_cells:
                    if not unique_cells or unique_cells[-1] != cell:
                        unique_cells.append(cell)
                extracted_lines.append(" | ".join(unique_cells))

    full_text = "\n".join(extracted_lines).strip()

    if not full_text:
        raise ValueError("DOCX document contains no readable text")

    return full_text
